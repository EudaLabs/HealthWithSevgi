#!/usr/bin/env python3
"""
Clean the filled ISO 42001 Report (Ch 1-3):
  - Remove template instruction paragraphs
  - Remove grading criteria / meta-info tables
  - Remove unfilled chapters 4-7, Appendix A & B
  - Trim TOC to Ch. 1-3 only
  - Clean up trailing empty paragraphs

Input:  ISO42001_Report_Ch1-3_FILLED.docx
Output: ISO42001_Report_Ch1-3_CLEAN.docx
"""
from pathlib import Path
from docx import Document
from docx.oxml.ns import qn

INPUT  = Path(__file__).parent / "ISO42001_Report_Ch1-3_FILLED.docx"
OUTPUT = Path(__file__).parent / "ISO42001_Report_Ch1-3_CLEAN.docx"


# ── XML helpers ──────────────────────────────────────────────────────────

def elem_text(elem):
    """Extract all visible text from an XML element."""
    return "".join(t.text for t in elem.iter(qn("w:t")) if t.text).strip()


def first_cell_text(tbl_elem):
    """Text of the very first <w:tc> in a table element."""
    for tc in tbl_elem.iter(qn("w:tc")):
        return "".join(t.text for t in tc.iter(qn("w:t")) if t.text).strip()
    return ""


def tag_local(elem):
    """Return the local tag name without namespace."""
    t = elem.tag
    return t.split("}")[-1] if "}" in t else t


# ── What to remove ───────────────────────────────────────────────────────

# Template instruction paragraphs within Ch. 1-3 (prefix match)
REMOVE_PARA_PREFIXES = [
    # Ch. 1
    "This chapter establishes the foundation of your AI Management System",
    "Identify all parties affected by or interested in your AI system",
    # Ch. 2
    "This chapter presents your team\u2019s AI Policy",   # smart quote
    "This chapter presents your team's AI Policy",         # straight quote
    "Map your Scrum roles to ISO 42001 AI governance roles",
    # Ch. 3
    "This chapter documents how your team manages the data lifecycle",
    "Document all datasets used in your tool",
    "Your ML tool relies on external components",
    "Third-Party & Dependency Register Template",
]

# Grading / meta-info tables within Ch. 1-3 (first-cell prefix match)
REMOVE_TABLE_PREFIXES = [
    "Rapor Hakkinda",                  # submission meta-info box
    "\u2705 B\u00f6l\u00fcm 1",       # ✅ Bölüm 1 grading
    "\u2705 B\u00f6l\u00fcm 2",       # ✅ Bölüm 2 grading
    "Neden Bu Tablo Onemli",           # explanation box (3.4)
    "Bolum 3 Degerlendirme",           # Ch. 3 grading (ASCII)
    "B\u00f6l\u00fcm 3 De\u011ferlendirme",  # Ch. 3 grading (Turkish)
    "FINAL TESLIM BOLUMLERI",          # final-submission notice → also triggers "cut here"
]

# TOC rows to remove (first-cell exact match)
TOC_REMOVE_CELLS = {
    "ILK TESLIM (8 Nisan - Sprint 3 sonrasi)",
    "INSTRUCTOR GERI BILDIRIM (8-15 Nisan)",
    "FINAL TESLIM (5 Mayis)",
    "4", "5", "6", "7", "A", "B",
}


# ── Main ─────────────────────────────────────────────────────────────────

def main():
    doc = Document(str(INPUT))
    body = doc.element.body
    children = list(body)

    to_remove = []
    past_ch3 = False          # once True → remove everything remaining
    removed_summary = []      # for logging

    # ── Pass 1: mark elements for removal ────────────────────────────────
    for child in children:
        tag = tag_local(child)

        # After FINAL TESLIM, remove everything except <w:sectPr>
        if past_ch3:
            if tag != "sectPr":
                to_remove.append(child)
            continue

        if tag == "tbl":
            fc = first_cell_text(child)
            # Check if this is the boundary marker
            if fc.startswith("FINAL TESLIM BOLUMLERI"):
                to_remove.append(child)
                removed_summary.append(f"[T] FINAL TESLIM (boundary)")
                past_ch3 = True
                continue
            # Check grading / meta tables
            for prefix in REMOVE_TABLE_PREFIXES:
                if fc.startswith(prefix):
                    to_remove.append(child)
                    removed_summary.append(f"[T] {fc[:60]}")
                    break

        elif tag == "p":
            text = elem_text(child)
            for prefix in REMOVE_PARA_PREFIXES:
                if text.startswith(prefix):
                    to_remove.append(child)
                    removed_summary.append(f"[P] {text[:60]}")
                    break

    # ── Pass 2: remove marked elements ───────────────────────────────────
    for elem in to_remove:
        body.remove(elem)

    print(f"Removed {len(to_remove)} elements:")
    for s in removed_summary[:20]:
        print(f"  - {s}")
    if len(removed_summary) > 20:
        print(f"  ... and {len(removed_summary) - 20} more (Ch.4-7 + Appendices)")

    # ── Pass 3: trim trailing empty paragraphs ───────────────────────────
    trimmed = 0
    while True:
        children = list(body)
        if len(children) < 2:
            break
        last = children[-1]
        second_last = children[-2]
        # Keep the final sectPr; trim empty <w:p> before it
        if tag_local(last) == "sectPr" and tag_local(second_last) == "p":
            if not elem_text(second_last):
                body.remove(second_last)
                trimmed += 1
                continue
        break
    if trimmed:
        print(f"Trimmed {trimmed} trailing empty paragraphs")

    # ── Pass 4: clean TOC table ──────────────────────────────────────────
    clean_toc(doc)

    # ── Save ─────────────────────────────────────────────────────────────
    doc.save(str(OUTPUT))
    print(f"\nSaved → {OUTPUT}")
    print(f"Done.")


def clean_toc(doc):
    """Remove rows for Ch. 4-7 and Appendices from the TOC table."""
    for table in doc.tables:
        if table.cell(0, 0).text.strip() == "No":
            tbl_elem = table._tbl
            rows = tbl_elem.findall(qn("w:tr"))
            removed = 0
            for row in rows:
                # First cell text
                for tc in row.iter(qn("w:tc")):
                    cell_text = "".join(
                        t.text for t in tc.iter(qn("w:t")) if t.text
                    ).strip()
                    if cell_text in TOC_REMOVE_CELLS:
                        tbl_elem.remove(row)
                        removed += 1
                    break
            print(f"TOC: removed {removed} rows (kept header + Ch. 1-3)")
            return
    print("WARNING: TOC table not found")


if __name__ == "__main__":
    main()
